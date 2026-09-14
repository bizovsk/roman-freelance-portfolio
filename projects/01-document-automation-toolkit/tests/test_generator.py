from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook

from src.generate_documents import (
    InputSchemaError,
    generate,
    load_input,
    parse_metrics,
    validate_row,
)

ROOT = Path(__file__).resolve().parents[1]
BRAND = ROOT / "examples/brand.json"
CSV_INPUT = ROOT / "examples/records.csv"
XLSX_INPUT = ROOT / "examples/records.xlsx"


def test_csv_and_xlsx_generate_same_record_counts(tmp_path):
    csv_generated, csv_rejected = generate(
        CSV_INPUT,
        BRAND,
        tmp_path / "csv",
        output_format="both",
    )
    xlsx_generated, xlsx_rejected = generate(
        XLSX_INPUT,
        BRAND,
        tmp_path / "xlsx",
        sheet_name="Records",
        output_format="both",
    )
    assert len(csv_generated) == len(xlsx_generated) == 3
    assert len(csv_rejected) == len(xlsx_rejected) == 1


def test_outputs_include_docx_pdf_manifest_validation_and_summary(tmp_path):
    generated, rejected = generate(
        CSV_INPUT,
        BRAND,
        tmp_path,
        output_format="both",
    )
    assert len(generated) == 3
    assert len(rejected) == 1
    assert (tmp_path / "generation_manifest.csv").exists()
    assert (tmp_path / "validation_report.csv").exists()
    assert (tmp_path / "run_summary.json").exists()
    for item in generated:
        assert (tmp_path / item["docx_file"]).exists()
        assert (tmp_path / item["pdf_file"]).exists()


def test_generated_documents_are_native_docx(tmp_path):
    generated, _ = generate(
        CSV_INPUT,
        BRAND,
        tmp_path,
        output_format="docx",
    )
    for item in generated:
        path = tmp_path / item["docx_file"]
        doc = Document(path)
        assert doc.paragraphs
        assert doc.tables
        with ZipFile(path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
            assert "w:txbxContent" not in document_xml


def test_generated_pdfs_have_valid_header(tmp_path):
    generated, _ = generate(
        CSV_INPUT,
        BRAND,
        tmp_path,
        output_format="pdf",
    )
    for item in generated:
        data = (tmp_path / item["pdf_file"]).read_bytes()
        assert data.startswith(b"%PDF-")
        assert len(data) > 1000


def test_optional_risk_section_is_omitted(tmp_path):
    generated, _ = generate(
        CSV_INPUT,
        BRAND,
        tmp_path,
        output_format="docx",
    )
    target = next(
        tmp_path / item["docx_file"]
        for item in generated
        if item["record_id"] == "PRJ-003"
    )
    text = "\n".join(paragraph.text for paragraph in Document(target).paragraphs)
    assert "Risks / open points" not in text
    assert "Next steps" in text


def test_missing_required_field_is_rejected():
    row = {
        "record_id": "X",
        "document_title": "",
        "recipient": "Team",
        "owner": "Owner",
        "status": "Draft",
        "created_on": "2026-09-10",
        "summary": "Summary",
        "due_on": "",
    }
    errors = validate_row(row)
    assert "Missing required field: document_title" in errors


def test_invalid_date_and_due_date_order_are_reported():
    invalid = {
        "record_id": "X",
        "document_title": "Title",
        "recipient": "Team",
        "owner": "Owner",
        "status": "Draft",
        "created_on": "10/09/2026",
        "summary": "Summary",
        "due_on": "",
    }
    assert any("Invalid date" in error for error in validate_row(invalid))

    backwards = invalid | {
        "created_on": "2026-09-10",
        "due_on": "2026-09-09",
    }
    assert "due_on cannot be earlier than created_on" in validate_row(backwards)


def test_duplicate_record_ids_are_rejected(tmp_path):
    path = tmp_path / "duplicates.csv"
    path.write_text(
        "record_id,document_title,recipient,owner,status,created_on,due_on,summary,metrics,risks,next_steps\n"
        "DUP,One,Team,A,Draft,2026-09-01,,Summary,,,\n"
        "DUP,Two,Team,B,Draft,2026-09-02,,Summary,,,\n",
        encoding="utf-8",
    )
    generated, rejected = generate(
        path,
        BRAND,
        tmp_path / "out",
        output_format="docx",
    )
    assert generated == []
    assert len(rejected) == 2
    assert all(
        "Duplicate record_id: DUP" in item["errors"]
        for item in rejected
    )


def test_malformed_metrics_are_rejected(tmp_path):
    path = tmp_path / "bad_metrics.csv"
    path.write_text(
        "record_id,document_title,recipient,owner,status,created_on,due_on,summary,metrics,risks,next_steps\n"
        "X,Title,Team,A,Draft,2026-09-01,,Summary,Rows 20,,\n",
        encoding="utf-8",
    )
    generated, rejected = generate(
        path,
        BRAND,
        tmp_path / "out",
        output_format="docx",
    )
    assert generated == []
    assert "expected Label=Value" in rejected[0]["errors"]


def test_xlsx_requires_sheet_when_workbook_has_multiple_sheets(tmp_path):
    path = tmp_path / "multi.xlsx"
    workbook = Workbook()
    workbook.active.title = "One"
    workbook.create_sheet("Two")
    workbook.save(path)
    with pytest.raises(InputSchemaError, match="choose one explicitly"):
        load_input(path)


def test_missing_required_column_fails_fast(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text(
        "record_id,document_title\nX,Title\n",
        encoding="utf-8",
    )
    with pytest.raises(InputSchemaError, match="Missing required column"):
        load_input(path)


def test_metrics_parser_is_strict():
    assert parse_metrics("Rows=20;Errors=2") == [
        ("Rows", "20"),
        ("Errors", "2"),
    ]
    with pytest.raises(ValueError):
        parse_metrics("Rows 20")
