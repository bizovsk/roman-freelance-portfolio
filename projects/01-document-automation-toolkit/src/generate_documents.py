from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

REQUIRED_FIELDS = (
    "record_id",
    "document_title",
    "recipient",
    "owner",
    "status",
    "created_on",
    "summary",
)
DATE_FIELDS = ("created_on", "due_on")
SUPPORTED_INPUTS = {".csv", ".xlsx"}


class InputSchemaError(ValueError):
    """Raised when an input file cannot be interpreted safely."""


@dataclass(frozen=True)
class Brand:
    organisation: str
    tagline: str
    accent_hex: str
    primary_hex: str
    footer: str

    @property
    def accent(self) -> RGBColor:
        return RGBColor.from_string(self.accent_hex)

    @property
    def primary(self) -> RGBColor:
        return RGBColor.from_string(self.primary_hex)


@dataclass(frozen=True)
class LoadedInput:
    records: list[dict[str, str]]
    input_type: str
    sheet: str | None


def load_brand(path: Path) -> Brand:
    data = json.loads(path.read_text(encoding="utf-8"))
    return Brand(**data)


def slug(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9]+", "_", value.strip())
    return value.strip("_") or "document"


def normalize_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def validate_headers(headers: Iterable[str | None]) -> list[str]:
    normalized = [
        str(header).strip()
        for header in headers
        if header is not None and str(header).strip()
    ]
    missing = [field for field in REQUIRED_FIELDS if field not in normalized]
    if missing:
        raise InputSchemaError("Missing required column(s): " + ", ".join(missing))
    return normalized


def load_csv(path: Path) -> LoadedInput:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise InputSchemaError("CSV file has no header row")
        validate_headers(reader.fieldnames)
        records = []
        for row_number, row in enumerate(reader, start=2):
            normalized = {
                str(key).strip(): normalize_value(value)
                for key, value in row.items()
                if key is not None
            }
            if not any(normalized.values()):
                continue
            normalized["_source_row"] = str(row_number)
            records.append(normalized)
    return LoadedInput(records=records, input_type="csv", sheet=None)


def load_xlsx(path: Path, sheet_name: str | None) -> LoadedInput:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            if len(workbook.sheetnames) != 1:
                raise InputSchemaError(
                    "XLSX contains multiple sheets; choose one explicitly with --sheet. "
                    f"Available: {', '.join(workbook.sheetnames)}"
                )
            sheet_name = workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise InputSchemaError(
                f"Worksheet '{sheet_name}' was not found. "
                f"Available: {', '.join(workbook.sheetnames)}"
            )

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise InputSchemaError("XLSX worksheet is empty") from exc

        headers = [normalize_value(value) for value in header_row]
        validate_headers(headers)
        records: list[dict[str, str]] = []
        for row_number, values in enumerate(rows, start=2):
            normalized = {
                header: normalize_value(value)
                for header, value in zip(headers, values, strict=False)
                if header
            }
            if not any(normalized.values()):
                continue
            normalized["_source_row"] = str(row_number)
            records.append(normalized)
        return LoadedInput(records=records, input_type="xlsx", sheet=sheet_name)
    finally:
        workbook.close()


def load_input(path: Path, sheet_name: str | None = None) -> LoadedInput:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_INPUTS:
        raise InputSchemaError(
            f"Unsupported input format '{suffix}'. Supported formats: CSV, XLSX"
        )
    if suffix == ".csv":
        if sheet_name:
            raise InputSchemaError("--sheet can only be used with XLSX input")
        return load_csv(path)
    return load_xlsx(path, sheet_name)


def parse_metrics(value: str) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for item in filter(None, (part.strip() for part in value.split(";"))):
        if "=" not in item:
            raise ValueError(f"Invalid metric '{item}': expected Label=Value")
        label, metric = item.split("=", 1)
        if not label.strip() or not metric.strip():
            raise ValueError(f"Invalid metric '{item}': expected Label=Value")
        pairs.append((label.strip(), metric.strip()))
    return pairs


def parse_items(value: str, delimiter: str) -> list[str]:
    return [part.strip() for part in value.split(delimiter) if part.strip()]


def parse_iso_date(raw: str) -> date:
    return datetime.strptime(raw, "%Y-%m-%d").date()


def validate_row(
    row: dict[str, str], duplicate_ids: set[str] | None = None
) -> list[str]:
    errors = [
        f"Missing required field: {field}"
        for field in REQUIRED_FIELDS
        if not row.get(field, "").strip()
    ]

    parsed_dates: dict[str, date] = {}
    for field in DATE_FIELDS:
        raw = row.get(field, "").strip()
        if raw:
            try:
                parsed_dates[field] = parse_iso_date(raw)
            except ValueError:
                errors.append(f"Invalid date in {field}: expected YYYY-MM-DD")

    if "created_on" in parsed_dates and "due_on" in parsed_dates:
        if parsed_dates["due_on"] < parsed_dates["created_on"]:
            errors.append("due_on cannot be earlier than created_on")

    record_id = row.get("record_id", "").strip()
    if duplicate_ids and record_id in duplicate_ids:
        errors.append(f"Duplicate record_id: {record_id}")

    metrics = row.get("metrics", "").strip()
    if metrics:
        try:
            parse_metrics(metrics)
        except ValueError as exc:
            errors.append(str(exc))

    return errors


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure_document(doc: Document, brand: Brand) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    for name, size, colour in (
        ("Title", 24, brand.primary),
        ("Heading 1", 15, brand.primary),
        ("Heading 2", 11.5, brand.accent),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = colour

    existing = {style.name for style in doc.styles}
    for name, size, bold, colour in (
        ("Toolkit Label", 8.5, True, RGBColor(105, 115, 125)),
        ("Toolkit Note", 9, False, RGBColor(105, 115, 125)),
    ):
        style = (
            doc.styles[name]
            if name in existing
            else doc.styles.add_style(name, 1)
        )
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = colour

    header = section.header.paragraphs[0]
    header.text = f"{brand.organisation.upper()}  |  DOCUMENT AUTOMATION DEMO"
    header.style = doc.styles["Toolkit Label"]

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f"{brand.footer}  |  Page ")
    add_page_number(footer)


def add_metadata_table(doc: Document, row: dict[str, str]) -> None:
    pairs = [
        ("Record", row["record_id"]),
        ("Prepared for", row["recipient"]),
        ("Owner", row["owner"]),
        ("Status", row["status"]),
        ("Created", row["created_on"]),
    ]
    if row.get("due_on"):
        pairs.append(("Due", row["due_on"]))

    table = doc.add_table(rows=len(pairs), cols=2)
    table.style = "Table Grid"
    for idx, (label, value) in enumerate(pairs):
        table.cell(idx, 0).text = label
        table.cell(idx, 1).text = value
        for run in table.cell(idx, 0).paragraphs[0].runs:
            run.bold = True
    doc.add_paragraph()


def build_docx(row: dict[str, str], brand: Brand, output_path: Path) -> None:
    doc = Document()
    configure_document(doc, brand)

    title = doc.add_paragraph(style="Title")
    title.add_run(row["document_title"])
    subtitle = doc.add_paragraph(style="Toolkit Note")
    subtitle.add_run(brand.tagline)

    add_metadata_table(doc, row)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(row["summary"])

    metrics = (
        parse_metrics(row.get("metrics", ""))
        if row.get("metrics", "")
        else []
    )
    if metrics:
        doc.add_heading("Key figures", level=1)
        table = doc.add_table(rows=len(metrics) + 1, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        for col in range(2):
            for run in table.cell(0, col).paragraphs[0].runs:
                run.bold = True
        for index, (label, value) in enumerate(metrics, start=1):
            table.cell(index, 0).text = label
            table.cell(index, 1).text = value

    risks = parse_items(row.get("risks", ""), ";")
    if risks:
        doc.add_heading("Risks / open points", level=1)
        for risk in risks:
            doc.add_paragraph(risk, style="List Bullet")

    next_steps = parse_items(row.get("next_steps", ""), "|")
    if next_steps:
        doc.add_heading("Next steps", level=1)
        for item in next_steps:
            doc.add_paragraph(item, style="List Number")

    note = doc.add_paragraph(style="Toolkit Note")
    note.add_run(
        "Generated from structured input. Optional empty sections are omitted automatically."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def hex_colour(value: str) -> colors.Color:
    raw = value.strip().lstrip("#")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", raw):
        raise ValueError(f"Invalid hex colour: {value}")
    return colors.HexColor(f"#{raw}")


def build_pdf(row: dict[str, str], brand: Brand, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    primary = hex_colour(brand.primary_hex)
    accent = hex_colour(brand.accent_hex)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ToolkitTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=26,
        textColor=primary,
        spaceAfter=5 * mm,
    )
    heading_style = ParagraphStyle(
        "ToolkitHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=primary,
        spaceBefore=4 * mm,
        spaceAfter=2 * mm,
    )
    body_style = ParagraphStyle(
        "ToolkitBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=14,
        spaceAfter=2 * mm,
    )
    note_style = ParagraphStyle(
        "ToolkitNote",
        parent=body_style,
        fontSize=8,
        textColor=colors.HexColor("#69737D"),
    )

    def page_decor(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 8)
        canvas.setFillColor(colors.HexColor("#69737D"))
        canvas.drawString(
            20 * mm,
            A4[1] - 12 * mm,
            f"{brand.organisation.upper()}  |  DOCUMENT AUTOMATION DEMO",
        )
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(
            A4[0] - 20 * mm,
            10 * mm,
            f"{brand.footer}  |  Page {doc.page}",
        )
        canvas.restoreState()

    document = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=18 * mm,
        title=row["document_title"],
        author=brand.organisation,
    )

    story = [
        Paragraph(row["document_title"], title_style),
        Paragraph(brand.tagline, note_style),
        Spacer(1, 2 * mm),
    ]

    metadata = [
        ["Record", row["record_id"]],
        ["Prepared for", row["recipient"]],
        ["Owner", row["owner"]],
        ["Status", row["status"]],
        ["Created", row["created_on"]],
    ]
    if row.get("due_on"):
        metadata.append(["Due", row["due_on"]])
    meta_table = Table(metadata, colWidths=[38 * mm, 112 * mm], hAlign="LEFT")
    meta_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D1D9")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F5F7")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TEXTCOLOR", (0, 0), (0, -1), primary),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend([meta_table, Spacer(1, 3 * mm)])

    story.extend(
        [
            Paragraph("Summary", heading_style),
            Paragraph(row["summary"], body_style),
        ]
    )

    metrics = (
        parse_metrics(row.get("metrics", ""))
        if row.get("metrics", "")
        else []
    )
    if metrics:
        story.append(Paragraph("Key figures", heading_style))
        metric_data = [["Metric", "Value"], *[[label, value] for label, value in metrics]]
        metric_table = Table(
            metric_data,
            colWidths=[95 * mm, 55 * mm],
            hAlign="LEFT",
        )
        metric_table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D1D9")),
                    ("BACKGROUND", (0, 0), (-1, 0), accent),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(metric_table)

    risks = parse_items(row.get("risks", ""), ";")
    if risks:
        story.append(Paragraph("Risks / open points", heading_style))
        story.append(
            ListFlowable(
                [ListItem(Paragraph(item, body_style)) for item in risks],
                bulletType="bullet",
                leftIndent=14,
            )
        )

    next_steps = parse_items(row.get("next_steps", ""), "|")
    if next_steps:
        story.append(Paragraph("Next steps", heading_style))
        story.append(
            ListFlowable(
                [ListItem(Paragraph(item, body_style)) for item in next_steps],
                bulletType="1",
                leftIndent=18,
            )
        )

    story.extend(
        [
            Spacer(1, 4 * mm),
            Paragraph(
                "Generated from structured input. Optional empty sections are omitted automatically.",
                note_style,
            ),
        ]
    )

    document.build(story, onFirstPage=page_decor, onLaterPages=page_decor)


def write_csv(
    path: Path,
    fieldnames: list[str],
    rows: list[dict[str, str]],
) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def generate(
    input_path: Path,
    brand_json: Path,
    output_dir: Path,
    sheet_name: str | None = None,
    output_format: str = "both",
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    if output_format not in {"docx", "pdf", "both"}:
        raise ValueError("output_format must be one of: docx, pdf, both")

    brand = load_brand(brand_json)
    loaded = load_input(input_path, sheet_name)
    records = loaded.records
    ids = [
        row.get("record_id", "").strip()
        for row in records
        if row.get("record_id", "").strip()
    ]
    duplicate_ids = {
        record_id
        for record_id, count in Counter(ids).items()
        if count > 1
    }

    generated: list[dict[str, str]] = []
    rejected: list[dict[str, str]] = []
    output_dir.mkdir(parents=True, exist_ok=True)

    for row in records:
        errors = validate_row(row, duplicate_ids)
        if errors:
            rejected.append(
                {
                    "source_row": row.get("_source_row", ""),
                    "record_id": row.get("record_id", ""),
                    "errors": " | ".join(errors),
                }
            )
            continue

        stem = f"{slug(row['record_id'])}_{slug(row['document_title'])}"
        docx_file = ""
        pdf_file = ""
        if output_format in {"docx", "both"}:
            docx_file = f"{stem}.docx"
            build_docx(row, brand, output_dir / docx_file)
        if output_format in {"pdf", "both"}:
            pdf_file = f"{stem}.pdf"
            build_pdf(row, brand, output_dir / pdf_file)

        generated.append(
            {
                "record_id": row["record_id"],
                "docx_file": docx_file,
                "pdf_file": pdf_file,
                "status": "generated",
            }
        )

    write_csv(
        output_dir / "generation_manifest.csv",
        ["record_id", "docx_file", "pdf_file", "status"],
        generated,
    )
    write_csv(
        output_dir / "validation_report.csv",
        ["source_row", "record_id", "errors"],
        rejected,
    )

    summary = {
        "input_file": input_path.name,
        "input_type": loaded.input_type,
        "sheet": loaded.sheet,
        "requested_output_format": output_format,
        "records_seen": len(records),
        "records_generated": len(generated),
        "records_rejected": len(rejected),
        "brand": asdict(brand),
    }
    (output_dir / "run_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    return generated, rejected


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Validate CSV/XLSX data and generate reusable DOCX/PDF business documents."
        )
    )
    parser.add_argument(
        "--input",
        type=Path,
        required=True,
        help="Input .csv or .xlsx file",
    )
    parser.add_argument(
        "--brand",
        type=Path,
        default=Path("examples/brand.json"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("examples/generated"),
    )
    parser.add_argument("--sheet", help="Worksheet name for XLSX input")
    parser.add_argument(
        "--format",
        dest="output_format",
        choices=("docx", "pdf", "both"),
        default="both",
    )
    args = parser.parse_args()

    try:
        generated, rejected = generate(
            args.input,
            args.brand,
            args.output,
            sheet_name=args.sheet,
            output_format=args.output_format,
        )
    except (InputSchemaError, ValueError) as exc:
        parser.error(str(exc))

    print(
        f"Generated {len(generated)} record(s); "
        f"rejected {len(rejected)} record(s)."
    )


if __name__ == "__main__":
    main()
